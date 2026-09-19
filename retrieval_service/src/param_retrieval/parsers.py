"""File-format adapters. Add a parser to PARSERS without touching the worker or API."""
import csv
import json
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from .ports import Chunk

IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp"}


def field_key(value):
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


def clean(value, blocked):
    if isinstance(value, dict):
        return {str(k): clean(v, blocked) for k, v in value.items()
                if field_key(k) not in blocked and field_key(re.sub(r"_\d+$", "", str(k))) not in blocked}
    if isinstance(value, (list, tuple)):
        return [clean(v, blocked) for v in value]
    if isinstance(value, (datetime, date)):
        return value.isoformat()  # preserve timezone/naive distinction, do not guess
    if isinstance(value, Decimal):
        return str(value)
    return value


def table_row(values, columns, locator):
    # Duplicate/blank headings get positional suffixes instead of overwriting values.
    seen, result = set(), {}
    for index, value in enumerate(values):
        name = str(columns[index] or f"column_{index + 1}") if index < len(columns) else f"column_{index + 1}"
        if name in seen:
            name = f"{name}_{index + 1}"
        seen.add(name)
        result[name] = value
    return Chunk(locator, attributes=result)


def csv_rows(path):
    with path.open(encoding="utf-8-sig", newline="") as source:
        reader = csv.reader(source)
        columns = next(reader, [])
        for number, values in enumerate(reader, 2):
            if len(values) != len(columns):
                raise ValueError(f"CSV logical row {number} has a different column count")
            yield table_row(values, columns, f"row:{number}")


def xlsx_rows(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            columns = next(rows, ())
            for number, values in enumerate(rows, 2):
                if any(v is not None for v in values):
                    yield table_row(values, columns, f"sheet:{sheet.title}|row:{number}")
    finally:
        workbook.close()


def pdf_pages(path):
    from pypdf import PdfReader
    for number, page in enumerate(PdfReader(path).pages, 1):
        text = page.extract_text() or ""
        if not text.strip() and page.images:
            raise ValueError(f"PDF page {number} requires OCR; no text was silently discarded")
        if text.strip():
            yield Chunk(f"page:{number}", text)


def docx_blocks(path):
    from docx import Document
    from docx.table import Table
    for number, block in enumerate(Document(path).iter_inner_content(), 1):
        text = "\n".join(" | ".join(c.text for c in r.cells) for r in block.rows) if isinstance(block, Table) else block.text
        yield Chunk(f"block:{number}", text)


def text_lines(path):
    with path.open(encoding="utf-8-sig") as source:
        for number, line in enumerate(source, 1):
            yield Chunk(f"line:{number}", line)


def json_rows(path):
    import ijson
    with path.open("rb") as source:
        first = source.read(1)
        while first and first.isspace():
            first = source.read(1)
        source.seek(0)
        prefix = "item" if first == b"[" else ""
        for number, row in enumerate(ijson.items(source, prefix), 1):
            yield Chunk(f"record:{number}", attributes=row if isinstance(row, dict) else {"value": row})


def jsonl_rows(path):
    with path.open(encoding="utf-8-sig") as source:
        for number, line in enumerate(source, 1):
            if line.strip():
                value = json.loads(line)
                yield Chunk(f"line:{number}", attributes=value if isinstance(value, dict) else {"value": value})


PARSERS = {".csv": csv_rows, ".xlsx": xlsx_rows, ".pdf": pdf_pages,
           ".docx": docx_blocks, ".txt": text_lines, ".md": text_lines,
           ".json": json_rows, ".jsonl": jsonl_rows}


def parse(path: Path, suffix: str, settings):
    if suffix in IMAGE_SUFFIXES:
        yield Chunk("image", image=path)
        return
    parser = PARSERS.get(suffix)
    if parser is None:
        raise ValueError(f"Unsupported format: {suffix}")
    blocked = {field_key(name) for name in settings.blocked_fields}
    for record in parser(path):
        attributes = clean(record.attributes, blocked)
        text = record.text or json.dumps(attributes, ensure_ascii=False, allow_nan=False)
        if not text.strip() or (not record.text and not attributes):
            continue
        step = settings.chunk_chars - settings.chunk_overlap
        for start in range(0, len(text), step):
            # No clipping of long cells/pages; each part gets its own source locator.
            yield Chunk(f"{record.locator}|char:{start}", text[start:start + settings.chunk_chars], attributes)
            if start + settings.chunk_chars >= len(text):
                break
